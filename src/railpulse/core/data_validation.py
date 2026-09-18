"""Dataset-specific schema checks. No feature fitting, cleaning, or raw-file mutation."""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
import re

import numpy as np
import pandas as pd

DOOR_COLUMNS = ["Datetime", "Motor current(mA)", "Motor Voltage(10mV)", "Motor electrodynamic force",
                "Door opening time(.1s)", "Door closing time(.1s)", "Close command", "Open command",
                "DCSR", "DCSL", "DLSR", "DLSL", "Door Opened", "Door Locked", "Door is opening",
                "Door is closing", "Door leaf position"]
RAIL_COLUMNS = ["Rotating speed"] + [f"{kind} of bearing in position {position} of car {car}"
                                     for car in range(1, 9) for position in range(1, 9)
                                     for kind in ("Vibration", "Shock")]
MISSING_MARKERS = {"", "NA", "N/A", "NaN", "nan", "null", "NULL", "None", "#N/A"}


def read_header(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        header = next(csv.reader(handle), [])
    if not header or len(set(header)) != len(header):
        raise ValueError(f"Missing or duplicate CSV columns: {path}")
    return header


def labels(path, filenames, columns):
    if read_header(path) != columns:
        raise ValueError(f"Unexpected label columns in {path}; expected {columns}")
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    if frame[columns[0]].duplicated().any() or set(frame[columns[0]]) != set(filenames):
        raise ValueError(f"Label coverage must exactly match unique training filenames: {path}")
    return frame


def door_time(value):
    try:
        values = [int(x) for x in str(value).split("-")]
        if len(values) != 7 or not 0 <= values[-1] <= 999:
            raise ValueError()
        return datetime(*values[:6], microsecond=values[6] * 1000)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"Invalid Door timestamp: {value}") from exc


def validate_door(base):
    frames, records = {}, []
    for name in ("Train.csv", "Test.csv"):
        path = base / name
        if read_header(path) != DOOR_COLUMNS:
            raise ValueError(f"Door requires the documented 17 ordered columns: {path}")
        frame = pd.read_csv(path)
        if frame.empty or not np.isfinite(frame.iloc[:, 1:].to_numpy(dtype=float)).all():
            raise ValueError(f"Empty or non-finite Door telemetry: {path}")
        times = pd.DatetimeIndex([door_time(x) for x in frame.Datetime])
        if times.has_duplicates or not times.is_monotonic_increasing:
            raise ValueError(f"Door timestamps must increase uniquely: {path}")
        frames[name] = frame
        records.append({"file_id": name, "rows": len(frame), "columns": len(frame.columns)})
    path = base / "Train_Segments_Answer.csv"
    expected = ["segment_id", "start_time", "end_time", "operation", "status", "n_rows"]
    if read_header(path) != expected:
        raise ValueError("Unexpected Door segment-label schema")
    annotations = pd.read_csv(path)
    if annotations.empty or annotations.segment_id.isna().any() or annotations.segment_id.duplicated().any():
        raise ValueError("Door segments require unique nonempty IDs")
    if not set(annotations.status) <= {"Normal", "Abnormal resistance"} or not set(annotations.operation) <= {"Open", "Close"}:
        raise ValueError("Invalid Door label or operation")
    positions = {timestamp: i for i, timestamp in enumerate(frames["Train.csv"].Datetime)}
    previous_end = -1
    for row in annotations.itertuples():
        a, b = positions.get(row.start_time), positions.get(row.end_time)
        if a is None or b is None or a <= previous_end or a >= b or row.n_rows != b - a + 1:
            raise ValueError(f"Door segment boundaries/count/overlap invalid: {row.segment_id}")
        previous_end = b
    return {"recordings": records, "labelled_segments": len(annotations),
            "label_counts": annotations.status.value_counts().to_dict(), "warnings": []}


def validate_numeric_recording(path, subsystem):
    if subsystem == "rail":
        header = read_header(path)
        if header != RAIL_COLUMNS:
            raise ValueError(f"Rail requires the documented 129 ordered channels: {path}")
        expected_columns, expected_rows, header_arg = 129, 10000, 0
    else:
        expected_columns, expected_rows, header_arg = 1, None, None
    count = 0
    for chunk in pd.read_csv(path, header=header_arg, dtype=np.float64, chunksize=20000, skip_blank_lines=False):
        if chunk.shape[1] != expected_columns or not np.isfinite(chunk.to_numpy()).all():
            raise ValueError(f"Invalid {subsystem} dimensions or non-finite values: {path}")
        if subsystem == "rail" and not np.isin(chunk.iloc[:, 0], [0, 1]).all():
            raise ValueError(f"Rail speed signal must be binary: {path}")
        count += len(chunk)
    if count < 3 or expected_rows is not None and count != expected_rows:
        raise ValueError(f"Unexpected {subsystem} sample count {count}: {path}")
    return {"file_id": path.name, "rows": count, "columns": expected_columns}


def validate_acv_recording(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows, ()))
        if not header or None in header or len(set(header)) != len(header):
            raise ValueError(f"Missing or duplicate ACV header: {path}")
        cars, telemetry_indices = {}, []
        for index, column in enumerate(header):
            match = re.fullmatch(r"Car (\d{2}) - (.+)", str(column))
            if match:
                cars.setdefault(match[1], []).append(match[2])
                telemetry_indices.append(index)
        if len(cars) != 8:
            raise ValueError(f"ACV requires eight distinct two-digit car IDs: {path}")
        time_columns = [i for i, c in enumerate(header) if str(c).lower() in ("time", "timestamp", "datetime")]
        if len(time_columns) != 1:
            raise ValueError(f"Expected one explicit ACV timestamp column: {path}")
        time_index = time_columns[0]
        row_count = missing = 0
        timestamps = []
        for row in rows:
            row_count += 1
            if row[time_index] is None:
                raise ValueError(f"Missing ACV timestamp at row {row_count + 1}: {path}")
            timestamps.append(row[time_index])
            # The distributed workbooks also encode missing telemetry as text
            # such as "None". Count it without changing categorical values.
            for index in telemetry_indices:
                value = row[index]
                if value is None or isinstance(value, str) and value.strip() in MISSING_MARKERS:
                    missing += 1
                elif isinstance(value, float) and not np.isfinite(value):
                    if np.isnan(value):
                        missing += 1
                    else:
                        raise ValueError(f"Infinite ACV telemetry at row {row_count + 1}: {path}")
        if row_count == 0:
            raise ValueError(f"Empty ACV case: {path}")
        times = pd.DatetimeIndex(pd.to_datetime(timestamps, errors="raise"))
        if times.isna().any():
            raise ValueError(f"Invalid ACV timestamps: {path}")
        warnings = []
        duplicates = int(times.duplicated().sum())
        if duplicates:
            warnings.append(f"{duplicates} duplicate timestamps")
        if not times.is_monotonic_increasing:
            warnings.append("Timestamps are not sorted; original order preserved")
        if missing:
            warnings.append("Missing telemetry retained for subsystem-specific handling")
        return {"file_id": path.name, "rows": row_count, "columns": len(header),
                "car_ids": sorted(cars), "parameters_per_car": cars,
                "missing_telemetry_fraction": missing / (row_count * len(telemetry_indices)),
                "duplicate_timestamps": duplicates, "warnings": warnings}
    finally:
        workbook.close()


def validate_subsystem(root, key, spec, entries):
    base = Path(root) / "data" / spec["destination"]
    expected = {Path(entry["path"]).relative_to(Path("data") / spec["destination"]).as_posix()
                for entry in entries if entry["subsystem"] == key and entry["path"].startswith("data/")}
    actual = {p.relative_to(base).as_posix() for p in base.rglob("*") if p.suffix.lower() in (".csv", ".xlsx")}
    if actual != expected:
        raise ValueError(f"Unexpected or missing raw {key} files: extra={sorted(actual - expected)}, missing={sorted(expected - actual)}")
    if key == "door":
        return validate_door(base)
    extension = "*.xlsx" if key == "acv" else "*.csv"
    train, test = sorted((base / "Train").glob(extension)), sorted((base / "Test").glob(extension))
    if len(train) != spec["train_files"] or len(test) != spec["test_files"]:
        raise ValueError(f"{key} train/test inventory differs from configuration")
    if {p.name for p in train}.intersection(p.name for p in test):
        raise ValueError(f"Overlapping train/test filenames for {key}")
    source_train = {e["git_blob_sha1"] for e in entries if e["subsystem"] == key and "/Train/" in e["path"]}
    source_test = {e["git_blob_sha1"] for e in entries if e["subsystem"] == key and "/Test/" in e["path"]}
    if source_train.intersection(source_test):
        raise ValueError(f"Identical source recordings in train and test for {key}")
    records = []
    for path in train + test:
        record = validate_acv_recording(path) if key == "acv" else validate_numeric_recording(path, key)
        records.append({"split": path.parent.name, **record})
    target = "faulty_car" if key == "acv" else "label" if key == "rail" else "damage"
    annotated = labels(base / "Train_Labels.csv", [p.name for p in train], ["filename", target])
    if key == "acv":
        car_ids = {r["file_id"]: r["car_ids"] for r in records}
        if any(row.faulty_car not in car_ids[row.filename] for row in annotated.itertuples()):
            raise ValueError("An ACV label names a car absent from its case")
        label_summary = {"training_cases": len(annotated)}
    elif key == "rail":
        if not set(annotated.label) <= {"Normal", "Side I", "Side II"}:
            raise ValueError("Invalid Rail class label")
        label_summary = annotated.label.value_counts().to_dict()
    else:
        damage = pd.to_numeric(annotated.damage, errors="raise").to_numpy()
        if not np.isfinite(damage).all() or np.any(damage <= 0):
            raise ValueError("SHM labels must be finite positive damage values")
        if len({r["rows"] for r in records}) != 1:
            raise ValueError("Official SHM recordings should have equal sample counts")
        label_summary = {"minimum_damage": float(damage.min()), "maximum_damage": float(damage.max())}
    return {"recordings": records, "label_summary": label_summary,
            "train_files": len(train), "test_files": len(test),
            "warnings": [f"{r['file_id']}: {w}" for r in records for w in r.get("warnings", [])]}
